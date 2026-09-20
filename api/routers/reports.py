import io
import json
from datetime import date, datetime
from typing import Optional, List

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import JSONResponse, StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func

from db.database import get_db
from db.models.user import User
from db.models.sale import Order, OrderItem, OrderTypeEnum, PaymentMethodEnum
from api.deps import get_current_user

router = APIRouter()

# ── Helper para validar permisos y filtrar órdenes ──
def get_orders_query(
    db: Session,
    current_user: User,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    order_type: Optional[str] = None,
    payment_method: Optional[str] = None,
):
    if not current_user.role or current_user.role.name not in ["admin", "administrador", "encargado", "supervisor"]:
        raise HTTPException(status_code=403, detail="No tienes permisos para ver reportes.")

    query = db.query(Order)

    # Restringir a la sucursal del supervisor
    if current_user.role.name in ["encargado", "supervisor"]:
        if not current_user.branch_id:
            raise HTTPException(status_code=403, detail="Supervisor sin sucursal asignada.")
        query = query.filter(Order.branch_id == current_user.branch_id)

    # Filtros
    if start_date:
        query = query.filter(func.date(Order.created_at) >= start_date)
    if end_date:
        query = query.filter(func.date(Order.created_at) <= end_date)
    if order_type:
        query = query.filter(Order.order_type == order_type)
    if payment_method:
        query = query.filter(Order.payment_method == payment_method)

    # Solo exitosas
    query = query.filter(Order.is_paid == True)
    
    return query.order_by(Order.created_at.desc())


@router.get("/sales")
def get_sales_report_json(
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    order_type: Optional[str] = None,
    payment_method: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    query = get_orders_query(db, current_user, start_date, end_date, order_type, payment_method)
    orders = query.all()

    data = []
    total_revenue = 0
    for o in orders:
        total_revenue += o.total_amount
        data.append({
            "id": o.id,
            "date": o.created_at.isoformat(),
            "total_amount": o.total_amount,
            "order_type": o.order_type,
            "payment_method": o.payment_method,
            "branch": o.branch.name if o.branch else "Online",
            "customer_email": o.user.email if o.user else "N/A"
        })
    
    return {
        "total_revenue": total_revenue,
        "count": len(data),
        "data": data
    }


@router.get("/sales/export")
def export_sales_report(
    format: str = Query("pdf", pattern="^(pdf|excel)$"),
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    order_type: Optional[str] = None,
    payment_method: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    query = get_orders_query(db, current_user, start_date, end_date, order_type, payment_method)
    orders = query.all()

    if format == "excel":
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Reporte de Ventas"
        
        # Headers
        ws.append(["ID", "Fecha", "Monto Total", "Tipo Orden", "Método Pago", "Sucursal", "Cliente"])
        
        for o in orders:
            branch_name = o.branch.name if o.branch else "Online"
            customer = o.user.email if o.user else "N/A"
            ws.append([
                o.id, 
                o.created_at.strftime("%Y-%m-%d %H:%M"), 
                o.total_amount, 
                o.order_type.value if hasattr(o.order_type, 'value') else o.order_type, 
                o.payment_method.value if hasattr(o.payment_method, 'value') else o.payment_method, 
                branch_name, 
                customer
            ])
            
        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        
        return StreamingResponse(
            stream, 
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=reporte_ventas.xlsx"}
        )

    elif format == "pdf":
        from fpdf import FPDF
        
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        
        pdf.cell(200, 10, txt="Reporte de Ventas", new_x="LMARGIN", new_y="NEXT", align="C")
        pdf.cell(200, 10, txt=f"Fecha de emision: {datetime.now().strftime('%Y-%m-%d')}", new_x="LMARGIN", new_y="NEXT", align="C")
        
        # Cabeceras
        pdf.set_font("Helvetica", 'B', 10)
        pdf.cell(20, 10, "ID", 1)
        pdf.cell(40, 10, "Fecha", 1)
        pdf.cell(30, 10, "Monto", 1)
        pdf.cell(40, 10, "Tipo", 1)
        pdf.cell(60, 10, "Sucursal", 1, new_x="LMARGIN", new_y="NEXT")
        
        pdf.set_font("Helvetica", '', 10)
        total = 0
        for o in orders:
            total += o.total_amount
            branch_name = o.branch.name if o.branch else "Online"
            pdf.cell(20, 10, str(o.id), 1)
            pdf.cell(40, 10, o.created_at.strftime("%Y-%m-%d"), 1)
            pdf.cell(30, 10, f"${o.total_amount:.2f}", 1)
            pdf.cell(40, 10, str(o.order_type.value if hasattr(o.order_type, 'value') else o.order_type), 1)
            pdf.cell(60, 10, branch_name, 1, new_x="LMARGIN", new_y="NEXT")
            
        pdf.set_font("Helvetica", 'B', 12)
        pdf.cell(200, 10, txt=f"TOTAL RECAUDADO: ${total:.2f}", new_x="LMARGIN", new_y="NEXT", align="L")
        
        pdf_bytes = pdf.output()
        stream = io.BytesIO(pdf_bytes)
        
        return StreamingResponse(
            stream, 
            media_type="application/pdf",
            headers={"Content-Disposition": "attachment; filename=reporte_ventas.pdf"}
        )

from pydantic import BaseModel
from core.config import settings

class ReportChatRequest(BaseModel):
    message: str
    sales_data: list
    
@router.post("/ai-chat")
def ai_reports_chat(
    req: ReportChatRequest,
    current_user: User = Depends(get_current_user)
):
    # Validar permisos
    if not current_user.role or current_user.role.name not in ["admin", "administrador", "encargado", "supervisor"]:
        raise HTTPException(status_code=403, detail="No tienes permisos para usar el asistente.")
        
    try:
        from google import genai
        from google.genai import types
        
        # Usamos la nueva API key proporcionada por el usuario
        client = genai.Client(api_key=settings.GEMINI_API_KEY_3)
        
        system_instruction = (
            "Eres un analista financiero experto de FashionStore. "
            "Tu tarea es responder preguntas sobre el siguiente reporte de ventas de forma clara, concisa y profesional.\n\n"
            f"DATOS DE VENTAS:\n{json.dumps(req.sales_data, ensure_ascii=False)}"
        )
        prompt = f"{system_instruction}\n\nPregunta del usuario:\n{req.message}\n\nRespuesta:"
        
        response = client.models.generate_content(
            model='gemini-1.5-flash',
            contents=prompt,
        )
        
        return {"response": response.text}
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))

